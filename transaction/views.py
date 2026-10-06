from django.shortcuts import redirect, render
from django.http import Http404, HttpResponse
from django.conf import settings 
from cart.models import Cart
import pandas as pd
from .models import productTransaction, transaction
from datetime import datetime, timedelta
from django.contrib.auth.decorators import login_required
from django import forms

class DateSelector(forms.Form):
    start_date = forms.DateField(widget = forms.SelectDateWidget())
    end_date = forms.DateField(widget = forms.SelectDateWidget())


class printer:
    printer = None

    def printReceipt(printText,times=0,*args,**kwargs):
        try:
            if printer.printer:
                printer.printer.text(printText)
                printer.printer.text(f"\nPrint Time: {datetime.now():%Y-%m-%d %H:%M}\n\n\n")
                # printer.printer.print_and_feed(n=3)
        except Exception as e: 
            printer.connectPrinter()
            if times <3:
                printer.printReceipt(printText, times+1)

    def connectPrinter():
        try:
            # Imported lazily: USB receipt printing is unavailable on hosts
            # like Vercel, and a top-level import here would crash the whole
            # URLconf (urls.py imports this module) and take down every route.
            from escpos.printer import Usb
            printer.printer = Usb(eval(settings.PRINTER_VENDOR_ID),eval(settings.PRINTER_PRODUCT_ID))
        except Exception as e:
            print(e)
            printer.printer = None


@login_required(login_url="/user/login/")
def transactionReceipt(request,transNo):
    try:
        t = transaction.objects.get(transaction_id=transNo)
        import ast
        items = ast.literal_eval(t.products)
        for it in items:
            it['line_total'] = round(float(it['price']) * float(it['quantity']), 2)
        return render(request,'receiptView.html',context={
            'receipt': t.receipt, 'transNo': transNo,
            'txn': t, 'items': items,
            'store_name': settings.STORE_NAME, 'store_address': settings.STORE_ADDRESS,
            'store_phone': settings.STORE_PHONE,
        })
    except transaction.DoesNotExist:
        raise Http404("No Transactions Found!!!")

@login_required(login_url="/user/login/")
def transactionPrintReceipt(request,transNo):
    try:
        receipt = transaction.objects.get(transaction_id=transNo).receipt
        if printer.printer is None:
            printer.connectPrinter() 
            print("Connecting Printer")
        if printer.printer: 
            printer.printReceipt(receipt)
        return redirect(f'/transaction_receipt/{transNo}/')
    except Exception as e:
        print(e)
        return redirect('register')


@login_required(login_url="/user/login/")
def transactionView(request, transNo=None):
    end_date=datetime.now().date()
    start_date=datetime.now().date()-timedelta(7)
    form = DateSelector(initial = {'end_date':end_date, 'start_date':start_date})
    if request.method == "POST":
        form = DateSelector(request.POST)
        if form.is_valid():
            end_date= form.cleaned_data['end_date']
            start_date= form.cleaned_data['start_date']
    transactions = list(transaction.objects.filter(transaction_dt__date__range = (start_date,end_date)).order_by('-transaction_dt').values('transaction_dt', 'transaction_id','total_sale','payment_type'))
    return render(request, 'transactions.html',
        context={'transactions':transactions,
            'form':form,})


@login_required(login_url="/user/login/")
def returnsTransaction(request):
    Cart(request).returns()
    return redirect('register')


@login_required(login_url="/user/login/")
def suspendTransaction(request):
    if Cart(request).isNotEmpty():
        if "Cart_Sessions" in request.session.keys():
            request.session["Cart_Sessions"][datetime.now().strftime('%Y%m%d%H%M%S%f')] = request.session[settings.CART_SESSION_ID]
            request.session.modified = True
        else:
            request.session["Cart_Sessions"] = {}
            request.session["Cart_Sessions"][datetime.now().strftime('%Y%m%d%H%M%S%f')] = request.session[settings.CART_SESSION_ID] 
    return redirect("cart_clear")


@login_required(login_url="/user/login/")
def recallTransaction(request, recallTransNo = None):
    if Cart(request).isNotEmpty():
        return redirect("suspend_transaction")
    if recallTransNo:
        request.session[settings.CART_SESSION_ID] = request.session["Cart_Sessions"][recallTransNo]
        del request.session["Cart_Sessions"][recallTransNo]
        request.session.modified = True
    elif "Cart_Sessions" in request.session.keys() and len(request.session["Cart_Sessions"]):
        return render(request, "recallTransaction.html", context={"obj_rt": request.session["Cart_Sessions"].keys()})
    return redirect("register")


@login_required(login_url="/user/login/")
def endTransactionReceipt(request,transNo):
    try:
        if request.GET["type"]=="cash":
            change = float(request.GET["value"]) - float(request.GET["total"])
            change = f"""<table class="table text-white h3 p-0 m-0"> 
                            <tr> 
                                <td class="text-left pl-5"> Total : </td> 
                                <td class="text-right pr-5"> PKR {request.GET["total"]}</td> 
                            </tr> 
                            <tr> 
                                <td class="text-left pl-5"> Cash : </td> 
                                <td class="text-right pr-5"> PKR {request.GET["value"]}</td> 
                            </tr> 
                            <tr class="h1 badge-danger" >  
                                <td style="padding-top:15px"> Change : </td> 
                                <td style="padding-top:15px"> PKR {change*(-1):.2f}</td> 
                            </tr> 
                        </table>"""
        elif request.GET["type"]=="card":
            change = f"""<table class="table text-white h3 p-0 m-0"> 
                            <tr> 
                                <td class="text-left pl-5"> Total : </td> 
                                <td class="text-right pr-5"> PKR {request.GET["total"]}</td> 
                            </tr> 
                            <tr> 
                                <td class="text-left pl-5"> Card : </td> 
                                <td class="text-right pr-5"> {request.GET["value"]}</td> 
                            </tr> 
                            
                        </table>
                        <div class="h1 badge-danger p-3" >  
                                 CARD TRANSACTION 
                            </div> 
                            """
        obj = transaction.objects.get(transaction_id=transNo)
        return render(request,'endTransaction.html',context={'receipt':obj.receipt,'change':change})
    except transaction.DoesNotExist:
        raise Http404("No Transactions Found!!!")


@login_required(login_url="/user/login/")
def setDiscount(request, percent):
    """Set (or clear, with 0) a manual whole-cart discount percentage."""
    try:
        percent = round(float(percent), 2)
    except (TypeError, ValueError):
        percent = 0
    if percent < 0 or percent > 100:
        percent = 0
    request.session["Discount_Percent"] = percent
    request.session.modified = True
    return redirect("register")


@login_required(login_url="/user/login/")
def endTransaction(request,type,value):
    try:
        # Card Transactions
        cart = request.session[settings.CART_SESSION_ID]
        gross_total = round(pd.DataFrame(cart).T["line_total"].astype(float).sum(),2)
        discount_percent = float(request.session.get("Discount_Percent", 0) or 0)
        discount_amount = round(gross_total * discount_percent / 100, 2)
        total = round(gross_total - discount_amount, 2)
        if type == "card": # Card Transaction
            # EBT Transaction
            if value=="EBT": 
                return_transaction = addTransaction(request.user,"EBT",total,cart,total,
                    discount_percent=discount_percent, discount_amount=discount_amount,
                    gross_total=gross_total)
            # DEBIT/CREDIT Transaction
            elif value=="DEBIT_CREDIT": 
                return_transaction = addTransaction(request.user,"DEBIT/CREDIT",total,cart,total,
                    discount_percent=discount_percent, discount_amount=discount_amount,
                    gross_total=gross_total)
        elif type=="cash": # Cash Transaction
            value = round(float(value),2)
            if value>= total: 
                return_transaction = addTransaction(request.user,"CASH",total,cart,value,
                    discount_percent=discount_percent, discount_amount=discount_amount,
                    gross_total=gross_total)
        if return_transaction:
            Cart(request).clear()
            request.session["Discount_Percent"] = 0
            request.session.modified = True
            return redirect(f"/endTransaction/{return_transaction.transaction_id}/?type={type}&value={value}&total={total}")
        return redirect("register")
    except Exception as e:
        print(e,type,value,request.user)
        return redirect("register")


from django.contrib.auth import get_user_model
from django.db.models import Count, Sum
from django.utils import timezone
from io import BytesIO
import csv


@login_required(login_url="/user/login/")
def sold_items_report(request):
    """Line-items sold within a chosen date range (defaults to today).

    Reuses the existing ``DateSelector`` widget so the filter looks the same
    as the other report screens in the app.
    """
    today = datetime.now().date()
    start_date = end_date = today
    if request.GET.get('start') and request.GET.get('end'):
        try:
            start_date = datetime.strptime(request.GET['start'], '%Y-%m-%d').date()
            end_date = datetime.strptime(request.GET['end'], '%Y-%m-%d').date()
        except ValueError:
            start_date = end_date = today
    form = DateSelector(initial={'start_date': start_date, 'end_date': end_date})
    if request.method == "POST":
        form = DateSelector(request.POST)
        if form.is_valid():
            start_date = form.cleaned_data['start_date']
            end_date = form.cleaned_data['end_date']

    rows = list(productTransaction.objects
        .filter(transaction_date_time__date__range=(start_date, end_date))
        .order_by('-transaction_date_time')
        .values(
            'transaction_date_time', 'transaction_id', 'transaction_id_num',
            'barcode', 'name', 'department', 'sales_price', 'qty',
            'tax_category', 'tax_percentage', 'tax_amount',
            'deposit_category', 'deposit', 'deposit_amount',
            'payment_type', 'transaction__user_id',
        )
    )

    users = {u.id: u.get_username() for u in get_user_model().objects.filter(
        id__in={r['transaction__user_id'] for r in rows if r['transaction__user_id']})}

    for r in rows:
        r['line_total'] = (r['sales_price'] or 0) * (r['qty'] or 0)
        r['cashier'] = users.get(r['transaction__user_id'])

    # Money totals come from the stored transactions so receipt-level
    # discounts are reflected exactly as they were charged.
    totals = transaction.objects.filter(
        id__in={r['transaction_id'] for r in rows}).aggregate(
        receipts=Count('id'), revenue=Sum('total_sale'),
        discount=Sum('discount_amount'), tax=Sum('tax_total'),
        deposit=Sum('deposit_total'))

    if request.GET.get('export') in ('csv', 'xlsx'):
        return _export_sold_items(request.GET['export'], rows, totals, start_date, end_date)

    quick_ranges = [
        ('Today', today, today),
        ('Yesterday', today - timedelta(days=1), today - timedelta(days=1)),
        ('Last 7 days', today - timedelta(days=6), today),
        ('This month', today.replace(day=1), today),
    ]

    context = {
        'form': form,
        'rows': rows,
        'total_qty': sum(r['qty'] or 0 for r in rows),
        'gross_sales': sum(r['line_total'] for r in rows),
        'total_receipts': totals['receipts'] or 0,
        'total_revenue': totals['revenue'] or 0,
        'total_discount': totals['discount'] or 0,
        'total_tax': totals['tax'] or 0,
        'total_deposit': totals['deposit'] or 0,
        'quick_ranges': quick_ranges,
        'start_date': start_date,
        'end_date': end_date,
        'start_date_iso': start_date.isoformat(),
        'end_date_iso': end_date.isoformat(),
        'store_name': settings.STORE_NAME,
        'currency': 'PKR',
    }
    return render(request, 'sold_items_report.html', context=context)


EXPORT_COLUMNS = ['Date', 'Time', 'Receipt #', 'Barcode', 'Product', 'Department', 'Qty',
                  'Unit Price', 'Line Total', 'Tax', 'Deposit', 'Payment', 'Cashier']
MONEY_COLUMNS = {'Unit Price', 'Line Total', 'Tax', 'Deposit'}


def _safe_text(value):
    """Stop spreadsheet apps from treating user-entered text as a formula."""
    value = '' if value is None else str(value)
    return "'" + value if value[:1] in ('=', '+', '-', '@') else value


def _export_sold_items(fmt, rows, totals, start_date, end_date):
    """Sold Items report as a CSV or Excel download, built from the same rows
    and receipt totals the page shows so the figures always match."""
    data = []
    for r in rows:
        dt = timezone.localtime(r['transaction_date_time'])
        data.append([
            dt.date(), dt.strftime('%H:%M'), r['transaction_id_num'],
            _safe_text(r['barcode']), _safe_text(r['name']), _safe_text(r['department']),
            r['qty'] or 0, float(r['sales_price'] or 0), float(r['line_total'] or 0),
            float(r['tax_amount'] or 0), float(r['deposit_amount'] or 0),
            r['payment_type'], _safe_text(r['cashier']),
        ])

    discount = float(totals['discount'] or 0)
    summary = [
        ['Total', '', '', '', '', '', sum(r[6] for r in data), '', sum(r[8] for r in data),
         sum(r[9] for r in data), sum(r[10] for r in data), '', ''],
        ['Less: receipt discounts', '', '', '', '', '', '', '', -discount, '', '', '', ''],
        ['Net revenue (incl. tax & deposit)', '', '', '', '', '', '', '', float(totals['revenue'] or 0),
         '', '', '', ''],
    ]
    if not discount:
        del summary[1]

    period = f"{start_date:%Y-%m-%d}" + ('' if start_date == end_date else f"_{end_date:%Y-%m-%d}")
    filename = f"sold-items_{period}.{fmt}"

    if fmt == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        response.write('\ufeff')  # BOM so Excel reads UTF-8 correctly
        writer = csv.writer(response)
        writer.writerow(EXPORT_COLUMNS)
        writer.writerows([[c.isoformat() if hasattr(c, 'isoformat') else c for c in row] for row in data])
        writer.writerow([])
        writer.writerows(summary)
        return response

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Sold Items'
    period_label = f"{start_date:%d %b %Y}" + ('' if start_date == end_date else f" – {end_date:%d %b %Y}")
    ws.append([f"{settings.STORE_NAME} — Sold Items, {period_label}"])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(EXPORT_COLUMNS)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='4E5AE8')
        cell.alignment = Alignment(horizontal='center')
    for row in data:
        ws.append(row)
    last_data_row = ws.max_row
    ws.append([])
    for row in summary:
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)

    for idx, name in enumerate(EXPORT_COLUMNS, start=1):
        letter = get_column_letter(idx)
        for (cell,) in ws.iter_rows(min_row=header_row + 1, min_col=idx, max_col=idx):
            if name in MONEY_COLUMNS:
                cell.number_format = '#,##0.00'
            elif name == 'Date' and cell.row <= last_data_row:
                cell.number_format = 'DD-MM-YYYY'
        widths = {'Date': 12, 'Time': 8, 'Receipt #': 22, 'Product': 32, 'Department': 16, 'Qty': 7}
        ws.column_dimensions[letter].width = widths.get(name, 14)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if data:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(EXPORT_COLUMNS))}{last_data_row}"

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def addTransaction(user,payment_type,total,cart,value,discount_percent=0,discount_amount=0.0,gross_total=None):
    transaction_id = datetime.now().strftime('%Y%m%d%H%M%S%f')
    short_id = transaction_id[:14]
    cart_df = pd.DataFrame(cart).T.reset_index(drop=True)
    cart_df.index = cart_df.index + 1
    tax_total = round(cart_df["tax_value"].astype(float).sum(),2)
    deposit_total = round(cart_df["deposit_value"].astype(float).sum(),2)
    discount_percent = round(float(discount_percent or 0), 2)
    discount_amount = round(float(discount_amount or 0), 2)
    if gross_total is None:
        gross_total = round(total + discount_amount, 2)
    receipt_date = datetime.now().strftime('%d %b %Y  %I:%M %p')
    w = settings.RECEIPT_CHAR_COUNT

    # Item lines
    items_lines = []
    for idx, row in cart_df.iterrows():
        items_lines.append(f" {str(idx)+')':<3}{str(row['name'])[:20]}")
        items_lines.append(f"    {row['barcode']:<12} x{int(row['quantity']):<3} PKR {row['price']}")
        dep = float(row['deposit_value'])
        if dep > 0:
            items_lines.append(f"    Deposit:  PKR {dep:.2f}")
    cart_string = "\n".join(items_lines)

    receipt = f"{'='*w}\n"
    receipt += f"{settings.STORE_NAME.center(w)}\n"
    receipt += f"{settings.STORE_ADDRESS.center(w)}\n"
    if settings.STORE_PHONE:
        receipt += f"{'Ph: '+str(settings.STORE_PHONE).center(w-4)}\n"
    receipt += f"{'='*w}\n"
    receipt += f"{receipt_date.center(w)}\n"
    receipt += f"{'Receipt #'+short_id.center(w-8)}\n"
    receipt += f"{'-'*w}\n"
    receipt += f" {'#':<3}{'Item':<13}{'Qty':>4}  {'Price':>8}\n"
    receipt += f" {'-'*w}\n"
    receipt += cart_string + "\n"
    receipt += f" {'-'*w}\n"
    receipt += f" {'Subtotal:':<15}PKR {round(gross_total-tax_total,2):>8.2f}\n"
    if tax_total > 0:
        receipt += f" {'Tax:':<15}PKR {tax_total:>8.2f}\n"
    if deposit_total > 0:
        receipt += f" {'Deposit:':<15}PKR {deposit_total:>8.2f}\n"
    if discount_percent > 0:
        receipt += f" {f'Discount {discount_percent:g}%:':<15}PKR {round(discount_amount,2):>8.2f}\n"
    receipt += f" {'='*w}\n"
    receipt += f" {'TOTAL:':<15}PKR {round(total,2):>8.2f}\n"
    receipt += f" {'-'*w}\n"
    receipt += f" {str(payment_type):<15}PKR {round(value,2):>8.2f}\n"
    receipt += f" {'CHANGE:':<15}PKR {round(value-total,2):>8.2f}\n"
    receipt += f"{'='*w}\n"
    receipt += f"{settings.RECEIPT_FOOTER.center(w)}\n"
    receipt += f"{'='*w}\n"
    receipt += f"{('Trans ID: '+transaction_id).center(w)}\n"
    receipt += f"{'='*w}"
    
    ## IF CASH DRAWER Connected uncomment below
    # if printer.printer and settings.CASH_DRAWER: 
    #     try: printer.printer.cashdraw(2)
    #     except: pass

    #Saving Transaction into Database
    return transaction.objects.create( transaction_id = transaction_id , transaction_dt = datetime.strptime(transaction_id[:-6],'%Y%m%d%H%M%S'),
            user = user, total_sale= total, sub_total = round(gross_total-tax_total,2),tax_total=tax_total, deposit_total = deposit_total,
            discount_percent = discount_percent if discount_percent > 0 else None,
            discount_amount = discount_amount if discount_percent > 0 else None,
            payment_type = payment_type, receipt = receipt, products = str(cart_df.to_dict('records')),
        )
